import io
import os
import pandas as pd
from PIL import Image
import requests
import streamlit as st
import streamlit.components.v1 as components
from pptx import Presentation
from pptx.util import Inches, Pt

# Importación para PDF real
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet

# Configuración de la aplicación
st.set_page_config(
    page_title="PARISS SPORT - Creador de Catálogos",
    page_icon="👟",
    layout="wide",
    initial_sidebar_state="collapsed",
)

URL_LOGO_PARISS = "https://lh3.googleusercontent.com/NsAjoLhqG1BA7hxWsjdAtEhT60MT6LeHE8dHutHOGS-ecU6z80pqmHm31C0_JvgvkwhMHAl-cgNJSZON4g=w1901"

# Estilos CSS Personalizados
st.markdown(
    """
    <style>
    .stApp {
        background-color: #FFFFFF;
        color: #111111;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }
    
    .stFileUploader label {
        font-weight: bold !important;
        font-size: 16px !important;
        color: #111111 !important;
    }
    div[data-testid="stFileUploader"] {
        background-color: #E5E7EB !important;
        border: 2px solid #9CA3AF !important;
        border-radius: 8px !important;
        padding: 10px !important;
        margin-bottom: 12px !important;
    }
    div[data-testid="stFileUploader"] * {
        color: #111111 !important;
    }

    div.stButton > button, div[data-testid="stPopover"] > button, div.stDownloadButton > button {
        background-color: #E5E7EB !important;
        color: #111111 !important;
        border: 2px solid #9CA3AF !important;
        font-weight: bold !important;
        font-size: 15px !important;
        border-radius: 6px !important;
        padding: 8px 16px !important;
        box-shadow: 0 2px 4px rgba(0,0,0,0.05) !important;
    }

    div.stButton > button:hover, div[data-testid="stPopover"] > button:hover, div.stDownloadButton > button:hover {
        background-color: #D1D5DB !important;
        border-color: #6B7280 !important;
        color: #000000 !important;
    }
    
    .tag-box {
        border: 1px solid #CCCCCC;
        background-color: #FFFFFF;
        padding: 4px 10px;
        font-size: 14px;
        font-weight: 600;
        display: inline-block;
    }
    </style>
""",
    unsafe_allow_html=True,
)

# Control de Navegación de Páginas
if "pagina_actual" not in st.session_state:
    st.session_state["pagina_actual"] = "principal"

# Capturar coordenadas guardadas desde el lienzo interactivo vía query_params
query_params = st.query_params
if "saved" in query_params:
    try:
        st.session_state["pos_l1_x"] = float(query_params.get("l1_x", 30))
        st.session_state["pos_l1_y"] = float(query_params.get("l1_y", 60))
        st.session_state["pos_l2_x"] = float(query_params.get("l2_x", 160))
        st.session_state["pos_l2_y"] = float(query_params.get("l2_y", 60))
        st.session_state["pos_title_x"] = float(query_params.get("title_x", 540))
        st.session_state["pos_title_y"] = float(query_params.get("title_y", 60))
        st.session_state["pos_code_x"] = float(query_params.get("code_x", 690))
        st.session_state["pos_code_y"] = float(query_params.get("code_y", 110))
        st.session_state["pos_img_x"] = float(query_params.get("img_x", 280))
        st.session_state["pos_img_y"] = float(query_params.get("img_y", 160))
        st.session_state["pos_img_w"] = float(query_params.get("img_w", 320))
        st.session_state["pos_img_h"] = float(query_params.get("img_h", 210))
        st.session_state["pos_det_x"] = float(query_params.get("det_x", 30))
        st.session_state["pos_det_y"] = float(query_params.get("det_y", 480))
        st.session_state["pos_prec_x"] = float(query_params.get("prec_x", 560))
        st.session_state["pos_prec_y"] = float(query_params.get("prec_y", 460))
        
        st.success("¡Diseño y posiciones guardadas correctamente!")
        st.query_params.clear()
        st.session_state["pagina_actual"] = "principal"
    except Exception as e:
        st.error(f"Error al guardar posiciones: {e}")

# Valores iniciales de posición en el lienzo
if "pos_l1_x" not in st.session_state: st.session_state["pos_l1_x"] = 30
if "pos_l1_y" not in st.session_state: st.session_state["pos_l1_y"] = 60
if "pos_l2_x" not in st.session_state: st.session_state["pos_l2_x"] = 160
if "pos_l2_y" not in st.session_state: st.session_state["pos_l2_y"] = 60
if "pos_title_x" not in st.session_state: st.session_state["pos_title_x"] = 540
if "pos_title_y" not in st.session_state: st.session_state["pos_title_y"] = 60
if "pos_code_x" not in st.session_state: st.session_state["pos_code_x"] = 690
if "pos_code_y" not in st.session_state: st.session_state["pos_code_y"] = 110
if "pos_img_x" not in st.session_state: st.session_state["pos_img_x"] = 280
if "pos_img_y" not in st.session_state: st.session_state["pos_img_y"] = 160
if "pos_img_w" not in st.session_state: st.session_state["pos_img_w"] = 320
if "pos_img_h" not in st.session_state: st.session_state["pos_img_h"] = 210
if "pos_det_x" not in st.session_state: st.session_state["pos_det_x"] = 30
if "pos_det_y" not in st.session_state: st.session_state["pos_det_y"] = 480
if "pos_prec_x" not in st.session_state: st.session_state["pos_prec_x"] = 560
if "pos_prec_y" not in st.session_state: st.session_state["pos_prec_y"] = 460

def formatear_texto(val):
    if pd.isna(val) or val is None:
        return ""
    val_str = str(val).strip()
    if val_str.lower() in ["nan", "none", "null"]:
        return ""
    return val_str

# ---------------------------------------------------------
# PANTALLA A: EDITOR INTERACTIVO ARRASTRABLE (DRAG & DROP)
# ---------------------------------------------------------
if st.session_state["pagina_actual"] == "editor":
    st.markdown("## 🖱️ Editor Visual Interactivo")
    st.markdown("• **Mover:** Haz clic al centro de cualquier cuadro y arrástralo.\n• **Agrandar / Achicar:** Toca el circulito azul de la esquina inferior derecha.\n• Haz clic en el botón verde **GUARDAR CAMBIOS** de la barra superior.")

    if st.button("⬅️ Volver sin guardar"):
        st.session_state["pagina_actual"] = "principal"
        st.rerun()

    st.markdown("<hr>", unsafe_allow_html=True)

    canvas_html = """
    <!DOCTYPE html>
    <html translate="no">
    <head>
    <meta name="google" content="notranslate" />
    <style>
      #canvas-container {
        width: 100%;
        max-width: 920px;
        height: 600px;
        border: 3px solid #111111;
        background-color: #FFFFFF;
        position: relative;
        overflow: hidden;
        margin: 0 auto;
        border-radius: 8px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.15);
        font-family: 'Segoe UI', Arial, sans-serif;
        user-select: none;
        padding-top: 45px;
        box-sizing: border-box;
      }
      #toolbar {
        position: absolute;
        top: 0;
        left: 0;
        right: 0;
        height: 45px;
        background: #F3F4F6;
        border-bottom: 2px solid #D1D5DB;
        display: flex;
        align-items: center;
        justify-content: space-between;
        padding: 0 15px;
        z-index: 2000;
      }
      .toolbar-title {
        font-weight: bold;
        font-size: 14px;
        color: #374151;
      }
      #save-btn {
        background-color: #16A34A;
        color: white;
        border: none;
        padding: 6px 16px;
        font-weight: bold;
        font-size: 14px;
        border-radius: 4px;
        cursor: pointer;
        box-shadow: 0 2px 4px rgba(0,0,0,0.1);
      }
      #save-btn:hover {
        background-color: #15803D;
      }
      .draggable {
        position: absolute;
        cursor: move;
        padding: 8px 12px;
        background: #FFFFFF;
        border: 2px dashed #2563EB;
        border-radius: 6px;
        box-shadow: 0 2px 6px rgba(0,0,0,0.1);
        display: inline-block;
        box-sizing: border-box;
      }
      .draggable:active {
        border-color: #16A34A;
        z-index: 1000;
      }
      .resizer {
        width: 12px;
        height: 12px;
        background: #2563EB;
        border: 2px solid #FFFFFF;
        border-radius: 50%;
        position: absolute;
        right: -6px;
        bottom: -6px;
        cursor: se-resize;
      }
      .title-box { font-weight: bold; font-size: 20px; color: #111; border-color: #000; }
      .code-box { font-weight: bold; font-size: 15px; color: #4B5563; }
      .img-box { background: #EFF6FF; color: #1D4ED8; font-weight: bold; text-align: center; display: flex; align-items: center; justify-content: center; border: 2px dashed #2563EB; }
      .price-box { font-weight: bold; font-size: 15px; background: #F9FAFB; border: 1px solid #D1D5DB; }
      .details-box { font-weight: bold; font-size: 13px; background: #F3F4F6; border: 1px solid #D1D5DB; }
      .logo-box { font-size: 11px; font-weight: bold; width: 110px; text-align: center; background: #F9FAFB; border: 2px dashed #4B5563; }
    </style>
    </head>
    <body translate="no">

    <div id="canvas-container" translate="no">
      <div id="toolbar">
        <span class="toolbar-title">🎨 Lienzo de Diseño Interactivo</span>
        <button id="save-btn" onclick="saveChanges()">💾 GUARDAR CAMBIOS</button>
      </div>

      <div id="item-logo1" class="draggable logo-box" style="left: __L1_X__px; top: __L1_Y__px;">
        LOGOTIPO MARCA
        <div class="resizer"></div>
      </div>

      <div id="item-logo2" class="draggable logo-box" style="left: __L2_X__px; top: __L2_Y__px;">
        LOGOTIPO EMPRESA
        <div class="resizer"></div>
      </div>
      
      <div id="item-title" class="draggable title-box" style="left: __TITLE_X__px; top: __TITLE_Y__px;">
        Calzado Adidas Runfalcon 6
        <div class="resizer"></div>
      </div>

      <div id="item-code" class="draggable code-box" style="left: __CODE_X__px; top: __CODE_Y__px;">
        CÓDIGO: IH1825
        <div class="resizer"></div>
      </div>
      
      <div id="item-img" class="draggable img-box" style="left: __IMG_X__px; top: __IMG_Y__px; width: __IMG_W__px; height: __IMG_H__px;">
        Imagen de la Zapatilla
        <div class="resizer"></div>
      </div>
      
      <div id="item-details" class="draggable details-box" style="left: __DET_X__px; top: __DET_Y__px;">
        MARCA: ADIDAS | GÉNERO: HOMBRE | CAT: ZAPATILLA
        <div class="resizer"></div>
      </div>
      
      <div id="item-price" class="draggable price-box" style="left: __PREC_X__px; top: __PREC_Y__px;">
        <div>PRECIO MAYORISTA: S/ 171.18</div>
        <div style="margin-top: 4px;">PRECIO RETAIL: S/ 229.00</div>
        <div class="resizer"></div>
      </div>
    </div>

    <script>
      const draggables = document.querySelectorAll('.draggable');
      const container = document.getElementById('canvas-container');

      draggables.forEach(draggable => {
        const resizer = draggable.querySelector('.resizer');

        draggable.addEventListener('mousedown', function(e) {
          if (e.target === resizer) return;

          let el = draggable;
          let startX = e.clientX - el.offsetLeft;
          let startY = e.clientY - el.offsetTop;

          function onMouseMove(e) {
            let left = e.clientX - startX;
            let top = e.clientY - startY;

            left = Math.max(0, Math.min(left, container.clientWidth - el.offsetWidth));
            top = Math.max(45, Math.min(top, container.clientHeight - el.offsetHeight));

            el.style.left = left + 'px';
            el.style.top = top + 'px';
          }

          function onMouseUp() {
            document.removeEventListener('mousemove', onMouseMove);
            document.removeEventListener('mouseup', onMouseUp);
          }

          document.addEventListener('mousemove', onMouseMove);
          document.addEventListener('mouseup', onMouseUp);
        });

        resizer.addEventListener('mousedown', function(e) {
          e.stopPropagation();
          let el = draggable;
          let startWidth = parseInt(document.defaultView.getComputedStyle(el).width, 10);
          let startHeight = parseInt(document.defaultView.getComputedStyle(el).height, 10);
          let startX = e.clientX;
          let startY = e.clientY;

          function doResize(e) {
            let newWidth = startWidth + e.clientX - startX;
            let newHeight = startHeight + e.clientY - startY;

            if (newWidth > 60) el.style.width = newWidth + 'px';
            if (newHeight > 30) el.style.height = newHeight + 'px';
          }

          function stopResize() {
            document.removeEventListener('mousemove', doResize);
            document.removeEventListener('mouseup', stopResize);
          }

          document.addEventListener('mousemove', doResize);
          document.addEventListener('mouseup', stopResize);
        });
      });

      function saveChanges() {
        const l1 = document.getElementById('item-logo1');
        const l2 = document.getElementById('item-logo2');
        const title = document.getElementById('item-title');
        const code = document.getElementById('item-code');
        const img = document.getElementById('item-img');
        const det = document.getElementById('item-details');
        const prec = document.getElementById('item-price');

        let url = window.parent.location.pathname + '?saved=true' +
          '&l1_x=' + parseInt(l1.style.left) + '&l1_y=' + parseInt(l1.style.top) +
          '&l2_x=' + parseInt(l2.style.left) + '&l2_y=' + parseInt(l2.style.top) +
          '&title_x=' + parseInt(title.style.left) + '&title_y=' + parseInt(title.style.top) +
          '&code_x=' + parseInt(code.style.left) + '&code_y=' + parseInt(code.style.top) +
          '&img_x=' + parseInt(img.style.left) + '&img_y=' + parseInt(img.style.top) + '&img_w=' + parseInt(img.style.width) + '&img_h=' + parseInt(img.style.height) +
          '&det_x=' + parseInt(det.style.left) + '&det_y=' + parseInt(det.style.top) +
          '&prec_x=' + parseInt(prec.style.left) + '&prec_y=' + parseInt(prec.style.top);

        window.parent.location.href = url;
      }
    </script>
    </body>
    </html>
    """

    canvas_html = (
        canvas_html.replace("__L1_X__", str(int(st.session_state["pos_l1_x"])))
        .replace("__L1_Y__", str(int(st.session_state["pos_l1_y"])))
        .replace("__L2_X__", str(int(st.session_state["pos_l2_x"])))
        .replace("__L2_Y__", str(int(st.session_state["pos_l2_y"])))
        .replace("__TITLE_X__", str(int(st.session_state["pos_title_x"])))
        .replace("__TITLE_Y__", str(int(st.session_state["pos_title_y"])))
        .replace("__CODE_X__", str(int(st.session_state["pos_code_x"])))
        .replace("__CODE_Y__", str(int(st.session_state["pos_code_y"])))
        .replace("__IMG_X__", str(int(st.session_state["pos_img_x"])))
        .replace("__IMG_Y__", str(int(st.session_state["pos_img_y"])))
        .replace("__IMG_W__", str(int(st.session_state["pos_img_w"])))
        .replace("__IMG_H__", str(int(st.session_state["pos_img_h"])))
        .replace("__DET_X__", str(int(st.session_state["pos_det_x"])))
        .replace("__DET_Y__", str(int(st.session_state["pos_det_y"])))
        .replace("__PREC_X__", str(int(st.session_state["pos_prec_x"])))
        .replace("__PREC_Y__", str(int(st.session_state["pos_prec_y"])))
    )

    components.html(canvas_html, height=630)

# ---------------------------------------------------------
# PANTALLA B: GENERADOR DE CATÁLOGOS PRINCIPAL
# ---------------------------------------------------------
else:
    head_c1, head_c2 = st.columns([2, 4])

    with head_c1:
        st.image(URL_LOGO_PARISS, width=180)

    st.markdown("<br><b>Campos informativos a incluir:</b>", unsafe_allow_html=True)
    fc1, fc2, fc3, fc4 = st.columns(4)

    with fc1:
        ver_codigo = st.checkbox("CÓDIGO", value=True)
        ver_nombre = st.checkbox("NOMBRE", value=True)
    with fc2:
        ver_costo = st.checkbox("COSTO", value=False)
        ver_mayorista = st.checkbox("PRECIO MAYORISTA", value=True)
    with fc3:
        ver_retail = st.checkbox("PRECIO RETAIL", value=True)
        ver_genero = st.checkbox("GÉNERO", value=True)
    with fc4:
        ver_categoria = st.checkbox("CATEGORÍA", value=True)
        ver_marca = st.checkbox("MARCA", value=True)

    st.markdown("<hr style='margin:15px 0;'>", unsafe_allow_html=True)

    col_izq, col_der = st.columns([1.1, 2.7])

    with col_izq:
        st.markdown("### PANEL DE ENTRADAS")

        excel_file = st.file_uploader("1. Excel de tus datos", type=["xlsx", "xls", "csv"], key="up_excel")
        imagenes_files = st.file_uploader("2. Carpeta de imágenes", type=["png", "jpg", "jpeg", "webp", "avif"], accept_multiple_files=True, key="up_imgs")
        logo_marca_file = st.file_uploader("3. Imagen de logo de marca", type=["png", "jpg", "jpeg", "webp", "avif"], key="up_lmarca")
        logo_empresa_file = st.file_uploader("4. Imagen de logo de empresa", type=["png", "jpg", "jpeg", "webp", "avif"], key="up_lemp")

    with col_der:
        top_pre1, top_pre2 = st.columns([1.8, 1.2])

        with top_pre1:
            st.markdown(
                """
            <div style="border: 2px solid black; background: white; padding: 8px 14px; font-weight: bold; font-size: 18px; display: inline-block;">
                Tu catálogo quedará <u>así</u> :
            </div>
            """,
                unsafe_allow_html=True,
            )

        with top_pre2:
            if st.button("⚙️ Abrir Editor Arrastrable", use_container_width=True):
                st.session_state["pagina_actual"] = "editor"
                st.rerun()

        st.write("")

        dict_imgs = {}
        if imagenes_files:
            for f in imagenes_files:
                name_clean = f.name.split(".")[0].strip().upper()
                try:
                    img_obj = Image.open(f).convert("RGB")
                    dict_imgs[name_clean] = img_obj
                except Exception:
                    pass

        df = None
        if excel_file is not None:
            if excel_file.name.endswith(".csv"):
                df = pd.read_csv(excel_file)
            else:
                df = pd.read_excel(excel_file)

            row = df.iloc[0]

            # Mapeo seguro de campos evitando errores NaN
            codigo = ""
            nombre = ""
            costo = ""
            mayorista = ""
            retail = ""
            genero = ""
            categoria = ""
            marca = ""

            for col in row.index:
                col_u = str(col).strip().upper()
                val_u = formatear_texto(row[col])
                
                if "COD" in col_u and not codigo: codigo = val_u
                elif "NOM" in col_u or "DESCRIP" in col_u or "PROD" in col_u and not nombre: nombre = val_u
                elif "COSTO" in col_u and not costo: costo = val_u
                elif "MAYORISTA" in col_u and not mayorista: mayorista = val_u
                elif ("RETAIL" in col_u or "PÚBLICO" in col_u or "PUBLICO" in col_u or "VENTA" in col_u) and not retail: retail = val_u
                elif "GENERO" in col_u or "GÉNERO" in col_u and not genero: genero = val_u
                elif "CAT" in col_u and not categoria: categoria = val_u
                elif "MARCA" in col_u and not marca: marca = val_u

            t_col1, t_col2, t_col3 = st.columns([1, 1, 1.8])

            with t_col1:
                if logo_marca_file:
                    st.image(logo_marca_file, width=110)
                else:
                    st.markdown("<b>[ LOGO MARCA ]</b>", unsafe_allow_html=True)

            with t_col2:
                if logo_empresa_file:
                    st.image(logo_empresa_file, width=110)
                else:
                    st.image(URL_LOGO_PARISS, width=110)

            with t_col3:
                if ver_nombre and nombre:
                    st.markdown(
                        f"<div style='text-align:right;'><span class='tag-box' style='font-size:18px; font-weight:bold;'><b>{nombre}</b></span></div>",
                        unsafe_allow_html=True,
                    )
                if ver_codigo and codigo:
                    st.markdown(
                        f"<div style='text-align:right; margin-top:5px;'><span class='tag-box' style='font-size:16px;'>CÓDIGO: {codigo}</span></div>",
                        unsafe_allow_html=True,
                    )

            st.write("")
            m_col1, m_col2, m_col3 = st.columns([0.1, 3, 0.1])
            with m_col2:
                if codigo in dict_imgs:
                    st.image(dict_imgs[codigo], width=320)
                else:
                    st.warning(f"📷 Imagen no encontrada ({codigo if codigo else 'Sin Código'})")

            b_col1, b_col2 = st.columns([2, 1])

            with b_col1:
                bot_list = []
                if ver_marca and marca:
                    bot_list.append(f"<span class='tag-box'>MARCA: {marca}</span>")
                if ver_genero and genero:
                    bot_list.append(f"<span class='tag-box'>GÉNERO: {genero}</span>")
                if ver_categoria and categoria:
                    bot_list.append(f"<span class='tag-box'>CATEGORÍA: {categoria}</span>")
                st.markdown(
                    f"<div style='margin-top:15px;'>{' &nbsp; '.join(bot_list)}</div>",
                    unsafe_allow_html=True,
                )

            with b_col2:
                if ver_costo and costo:
                    st.markdown(
                        f"<div style='text-align:right; margin-bottom:4px;'><span class='tag-box'>COSTO: S/ {costo}</span></div>",
                        unsafe_allow_html=True,
                    )
                if ver_mayorista and mayorista:
                    st.markdown(
                        f"<div style='text-align:right; margin-bottom:4px;'><span class='tag-box' style='font-weight:bold;'>PRECIO MAYORISTA: S/ {mayorista}</span></div>",
                        unsafe_allow_html=True,
                    )
                if ver_retail and retail:
                    st.markdown(
                        f"<div style='text-align:right;'><span class='tag-box' style='font-weight:bold;'>PRECIO RETAIL: S/ {retail}</span></div>",
                        unsafe_allow_html=True,
                    )

            # BOTONES DE EXPORTACIÓN Y DESCARGA
            st.markdown("<br><br>", unsafe_allow_html=True)
            e_col1, e_col2, e_col3 = st.columns(3)

            def generar_pptx_completo(es_a4=False):
                prs = Presentation()
                slide_w_in = 11.69 if es_a4 else 13.33
                slide_h_in = 8.27 if es_a4 else 7.5
                prs.slide_width = Inches(slide_w_in)
                prs.slide_height = Inches(slide_h_in)
                blank_layout = prs.slide_layouts[6]

                canvas_w_px = 920.0
                canvas_h_px = 600.0

                for _, r in df.iterrows():
                    slide = prs.slides.add_slide(blank_layout)
                    
                    c_cod, c_nom, c_mar, c_gen, c_cat, c_may, c_ret = "", "", "", "", "", "", ""
                    for col in r.index:
                        col_u = str(col).strip().upper()
                        val_u = formatear_texto(r[col])
                        if "COD" in col_u and not c_cod: c_cod = val_u
                        elif ("NOM" in col_u or "DESCRIP" in col_u or "PROD" in col_u) and not c_nom: c_nom = val_u
                        elif "MARCA" in col_u and not c_mar: c_mar = val_u
                        elif ("GENERO" in col_u or "GÉNERO" in col_u) and not c_gen: c_gen = val_u
                        elif "CAT" in col_u and not c_cat: c_cat = val_u
                        elif "MAYORISTA" in col_u and not c_may: c_may = val_u
                        elif ("RETAIL" in col_u or "PÚBLICO" in col_u or "PUBLICO" in col_u or "VENTA" in col_u) and not c_ret: c_ret = val_u

                    l1_x_in = (st.session_state["pos_l1_x"] / canvas_w_px) * slide_w_in
                    l1_y_in = (st.session_state["pos_l1_y"] / canvas_h_px) * slide_h_in

                    if logo_marca_file:
                        try:
                            l_m_bytes = io.BytesIO(logo_marca_file.getvalue())
                            slide.shapes.add_picture(l_m_bytes, Inches(l1_x_in), Inches(l1_y_in), width=Inches(1.5))
                        except Exception:
                            pass

                    l2_x_in = (st.session_state["pos_l2_x"] / canvas_w_px) * slide_w_in
                    l2_y_in = (st.session_state["pos_l2_y"] / canvas_h_px) * slide_h_in

                    if logo_empresa_file:
                        try:
                            l_e_bytes = io.BytesIO(logo_empresa_file.getvalue())
                            slide.shapes.add_picture(l_e_bytes, Inches(l2_x_in), Inches(l2_y_in), width=Inches(1.5))
                        except Exception:
                            pass
                    else:
                        try:
                            res = requests.get(URL_LOGO_PARISS)
                            if res.status_code == 200:
                                l_e_bytes = io.BytesIO(res.content)
                                slide.shapes.add_picture(l_e_bytes, Inches(l2_x_in), Inches(l2_y_in), width=Inches(1.5))
                        except Exception:
                            pass

                    t_x_in = (st.session_state["pos_title_x"] / canvas_w_px) * slide_w_in
                    t_y_in = (st.session_state["pos_title_y"] / canvas_h_px) * slide_h_in
                    if ver_nombre and c_nom:
                        tx = slide.shapes.add_textbox(Inches(t_x_in), Inches(t_y_in), Inches(5), Inches(1.0))
                        tf = tx.text_frame
                        p = tf.paragraphs[0]
                        p.text = c_nom
                        p.font.size = Pt(20)
                        p.font.bold = True

                    c_x_in = (st.session_state["pos_code_x"] / canvas_w_px) * slide_w_in
                    c_y_in = (st.session_state["pos_code_y"] / canvas_h_px) * slide_h_in
                    if ver_codigo and c_cod:
                        tx_c = slide.shapes.add_textbox(Inches(c_x_in), Inches(c_y_in), Inches(4), Inches(0.8))
                        tf_c = tx_c.text_frame
                        p_c = tf_c.paragraphs[0]
                        p_c.text = f"CÓDIGO: {c_cod}"
                        p_c.font.size = Pt(15)
                        p_c.font.bold = True

                    img_x_in = (st.session_state["pos_img_x"] / canvas_w_px) * slide_w_in
                    img_y_in = (st.session_state["pos_img_y"] / canvas_h_px) * slide_h_in
                    img_w_in = (st.session_state["pos_img_w"] / canvas_w_px) * slide_w_in
                    if c_cod in dict_imgs:
                        img_buf = io.BytesIO()
                        dict_imgs[c_cod].save(img_buf, format="PNG")
                        img_buf.seek(0)
                        slide.shapes.add_picture(img_buf, Inches(img_x_in), Inches(img_y_in), width=Inches(max(2.0, img_w_in)))

                    det_x_in = (st.session_state["pos_det_x"] / canvas_w_px) * slide_w_in
                    det_y_in = (st.session_state["pos_det_y"] / canvas_h_px) * slide_h_in
                    tx_det = slide.shapes.add_textbox(Inches(det_x_in), Inches(det_y_in), Inches(6), Inches(0.8))
                    tf_det = tx_det.text_frame
                    p_det = tf_det.paragraphs[0]
                    p_det.text = f"MARCA: {c_mar}  |  GÉNERO: {c_gen}  |  CAT: {c_cat}"
                    p_det.font.size = Pt(13)
                    p_det.font.bold = True

                    prec_x_in = (st.session_state["pos_prec_x"] / canvas_w_px) * slide_w_in
                    prec_y_in = (st.session_state["pos_prec_y"] / canvas_h_px) * slide_h_in
                    tx2 = slide.shapes.add_textbox(Inches(prec_x_in), Inches(prec_y_in), Inches(5), Inches(1.2))
                    tf2 = tx2.text_frame
                    
                    has_paragraph = False
                    if ver_mayorista and c_may:
                        p3 = tf2.paragraphs[0]
                        p3.text = f"PRECIO MAYORISTA: S/ {c_may}"
                        p3.font.size = Pt(16)
                        p3.font.bold = True
                        has_paragraph = True

                    if ver_retail and c_ret:
                        p4 = tf2.add_paragraph() if has_paragraph else tf2.paragraphs[0]
                        p4.text = f"PRECIO RETAIL: S/ {c_ret}"
                        p4.font.size = Pt(16)
                        p4.font.bold = True

                ppt_io = io.BytesIO()
                prs.save(ppt_io)
                ppt_io.seek(0)
                return ppt_io

            def generar_pdf_real():
                pdf_buffer = io.BytesIO()
                doc = SimpleDocTemplate(pdf_buffer, pagesize=landscape(letter), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
                story = []
                styles = getSampleStyleSheet()

                for _, r in df.iterrows():
                    c_cod, c_nom, c_mar, c_gen, c_cat, c_may, c_ret = "", "", "", "", "", "", ""
                    for col in r.index:
                        col_u = str(col).strip().upper()
                        val_u = formatear_texto(r[col])
                        if "COD" in col_u and not c_cod: c_cod = val_u
                        elif ("NOM" in col_u or "DESCRIP" in col_u or "PROD" in col_u) and not c_nom: c_nom = val_u
                        elif "MARCA" in col_u and not c_mar: c_mar = val_u
                        elif ("GENERO" in col_u or "GÉNERO" in col_u) and not c_gen: c_gen = val_u
                        elif "CAT" in col_u and not c_cat: c_cat = val_u
                        elif "MAYORISTA" in col_u and not c_may: c_may = val_u
                        elif ("RETAIL" in col_u or "PÚBLICO" in col_u or "PUBLICO" in col_u or "VENTA" in col_u) and not c_ret: c_ret = val_u

                    p_title = Paragraph(f"<b>{c_nom}</b><br/>CÓDIGO: {c_cod}", styles['Heading2'])
                    p_det = Paragraph(f"<b>MARCA:</b> {c_mar} &nbsp;&nbsp; <b>GÉNERO:</b> {c_gen} &nbsp;&nbsp; <b>CAT:</b> {c_cat}", styles['Normal'])
                    p_price = Paragraph(f"<b>P. MAYORISTA: S/ {c_may}</b><br/><b>P. RETAIL: S/ {c_ret}</b>", styles['Heading3'])

                    story.append(p_title)
                    story.append(Spacer(1, 10))

                    if c_cod in dict_imgs:
                        img_buf = io.BytesIO()
                        dict_imgs[c_cod].save(img_buf, format="PNG")
                        img_buf.seek(0)
                        rl_img = RLImage(img_buf, width=280, height=200)
                        story.append(rl_img)

                    story.append(Spacer(1, 15))
                    story.append(p_det)
                    story.append(Spacer(1, 10))
                    story.append(p_price)
                    story.append(Spacer(1, 40))

                doc.build(story)
                pdf_buffer.seek(0)
                return pdf_buffer

            with e_col1:
                ppt_16x9 = generar_pptx_completo(es_a4=False)
                st.download_button(
                    label="Exportar a ppt (16:9)",
                    data=ppt_16x9,
                    file_name="Catalogo_Pariss_Sport_16x9.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True,
                )

            with e_col2:
                pdf_bytes = generar_pdf_real()
                st.download_button(
                    label="Exportar a PDF",
                    data=pdf_bytes,
                    file_name="Catalogo_Pariss_Sport.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                )

            with e_col3:
                ppt_a4 = generar_pptx_completo(es_a4=True)
                st.download_button(
                    label="Exportar a ppt (A4 Imprimible)",
                    data=ppt_a4,
                    file_name="Catalogo_Pariss_Sport_A4.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True,
                )
        else:
            st.info("Sube tu archivo Excel en el panel de la izquierda para cargar los productos.")
            